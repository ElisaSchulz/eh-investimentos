"""
Carta mensal: gera os 3 gráficos do mês (índices americanos, mercados globais e
Ibovespa) e monta o Word a partir do modelo de cada marca (EH e HZ).

Rodar no Spyder: botão Run. Por padrão usa o mês anterior ao de hoje.
Para outro mês, preencha MES abaixo (ex.: "2026-09").

O modelo da marca fica em modelos/Carta_Mensal_Template_EH.docx (e _HZ.docx).
Os marcadores no modelo são {{MONTH}}, {{YEAR}}, {{PLOT1}}, {{PLOT2}} e {{PLOT3}}.
Marca sem modelo na pasta é pulada com um aviso.
"""
import os
import sys
from datetime import timedelta

sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))

import matplotlib

if os.environ.get("GITHUB_ACTIONS"):
    matplotlib.use("Agg")  # no GitHub não há tela; no Spyder os gráficos também aparecem no painel
import matplotlib.pyplot as plt
import numpy as np
import pandas as pd
from matplotlib.ticker import FuncFormatter

from comum import AQUI, fechamentos_yahoo, limites_do_mes, mes_referencia, nome_mes, pasta_saida

# ------------------------------------------------------------------ parâmetros
MES = None            # None = mês anterior. Ou "AAAA-MM".
PASTA_MODELOS = os.environ.get("PASTA_MODELOS") or os.path.join(AQUI, "modelos")

MARCAS = {
    "EH": {
        "modelo": "Carta_Mensal_Template_EH.docx",
        "nome_saida": "Carta Mensal EH",
        "cores": {"primary": "#002060", "secondary": "#ed9a1e", "neutral": "#BFBFBF"},
    },
    "HZ": {
        "modelo": "Carta_Mensal_Template_HZ.docx",
        "nome_saida": "Carta Mensal HZ",
        "cores": {"primary": "#1b0d6a", "secondary": "#649f98", "neutral": "#BFBFBF"},
    },
}

INDICES = {
    "^BVSP": "Ibovespa (Brasil)",
    "^GSPC": "S&P 500 (EUA)",
    "^IXIC": "Nasdaq (EUA)",
    "^DJI": "Dow Jones (EUA)",
    "^STOXX50E": "Euro Stoxx 50 (Europa)",
    "^FTSE": "FTSE 100 (Reino Unido)",
    "^GDAXI": "DAX (Alemanha)",
    "000001.SS": "SSE Composite (China)",
    "^HSCE": "Hang Seng China Enterprises (HK)",
    "^N225": "Nikkei 225 (Japão)",
}
EUA = ["^GSPC", "^IXIC", "^DJI"]

plt.rcParams["font.family"] = ["Arial", "Liberation Sans", "DejaVu Sans"]


# ------------------------------------------------------------------ dados
def retorno_no_mes(ticker, inicio, fim):
    """Fechamentos do mês precedidos do último fechamento do mês anterior (base = 0%).

    Usar o último pregão anterior como base evita perder o retorno do primeiro dia
    quando o último dia do mês anterior cai num fim de semana ou feriado.
    """
    s = fechamentos_yahoo(ticker, inicio - timedelta(days=15), fim + timedelta(days=1)).dropna()
    if s.empty:
        return pd.DataFrame(columns=["Close", "Cumulative Return"], index=pd.DatetimeIndex([]), dtype=float)
    antes, no_mes = s[s.index.date < inicio], s[s.index.date >= inicio]
    if no_mes.empty:
        return pd.DataFrame(columns=["Close", "Cumulative Return"], index=pd.DatetimeIndex([]), dtype=float)
    serie = pd.concat([antes.iloc[-1:], no_mes]) if not antes.empty else no_mes
    return pd.DataFrame({"Close": serie, "Cumulative Return": serie / serie.iloc[0] - 1})


def baixar_indices(inicio, fim):
    dados = {}
    for t in INDICES:
        d = retorno_no_mes(t, inicio, fim)
        if d.empty:
            print(f"Aviso: sem dados de {t} ({INDICES[t]}) no mês.")
        dados[t] = d
    df = pd.concat(dados, axis=1).sort_index()
    df.index = df.index.strftime("%d/%m/%Y")
    return df


def quatro_datas(indice):
    return [indice[0], indice[len(indice) // 3], indice[2 * len(indice) // 3], indice[-1]]


def limpar_eixos(ax, lados=("right", "left", "top", "bottom")):
    for lado in lados:
        ax.spines[lado].set_visible(False)
    ax.tick_params(axis="both", which="both", length=0)


# ------------------------------------------------------------------ gráficos
def grafico_eua(df, cores, caminho):
    d = df[EUA].dropna()
    fig, ax = plt.subplots(figsize=(12, 6))
    cor = {"^GSPC": cores["secondary"], "^IXIC": cores["primary"], "^DJI": cores["neutral"]}
    for t in EUA:
        serie = d[(t, "Cumulative Return")]
        ax.plot(d.index, serie, label=f"{INDICES[t]}: {serie.iloc[-1] * 100:.2f}%",
                color=cor[t], linewidth=3.5)
    ax.set_xticks(quatro_datas(d.index))
    ax.tick_params(axis="x", labelsize=12)
    limpar_eixos(ax)
    ax.axhline(0, color="gray", linewidth=1, alpha=0.7)
    ax.yaxis.set_major_formatter(FuncFormatter(lambda y, _: f"{y * 100:.0f}%"))
    ax.set_title("Índices Americanos - Retorno Acumulado", fontsize=18, fontweight="bold")
    ax.legend(frameon=False, fontsize=12)
    salvar(fig, caminho)


def grafico_global(df, cores, caminho):
    resumo = pd.Series({INDICES[t]: df[(t, "Cumulative Return")].dropna().iloc[-1]
                        for t in INDICES if not df[(t, "Cumulative Return")].dropna().empty})
    resumo = resumo.sort_values()
    fig, ax = plt.subplots(figsize=(12, 6))
    cor = [cores["secondary"] if nome == INDICES["^BVSP"] else cores["primary"] for nome in resumo.index]
    barras = ax.barh(resumo.index, resumo.values, color=cor, height=0.5)
    for barra, valor in zip(barras, resumo.values):
        x, ha = (barra.get_width() + 0.002, "left") if valor >= 0 else (barra.get_width() - 0.002, "right")
        ax.text(x, barra.get_y() + barra.get_height() / 2, f"{valor * 100:.2f}%",
                va="center", ha=ha, fontsize=12)
    ax.set_title("Mercados Globais - Retorno Acumulado no Mês", fontsize=18, fontweight="bold")
    ax.tick_params(axis="both", labelsize=12)
    ax.xaxis.set_major_formatter(FuncFormatter(lambda x, _: f"{x * 100:.0f}%"))
    ax.xaxis.set_major_locator(plt.MultipleLocator(0.05))
    limpar_eixos(ax, ("right", "left", "top"))
    ax.axvline(x=0, color="gray", linewidth=1.5, alpha=0.7)
    folga = 0.02
    ax.set_xlim(np.floor((resumo.min() - folga) * 100) / 100, np.ceil((resumo.max() + folga) * 100) / 100)
    salvar(fig, caminho)


def grafico_ibov(df, cores, caminho):
    d = df["^BVSP"].dropna()
    fig, ax1 = plt.subplots(figsize=(12, 6))
    ax1.set_zorder(3)
    ax1.patch.set_alpha(0)
    linha = ax1.plot(d.index, d["Close"], color=cores["primary"], linewidth=3.5, zorder=3)
    ax1.set_ylabel("Pontos Ibovespa", fontsize=12)

    ax2 = ax1.twinx()
    ax2.set_zorder(2)
    ax2.patch.set_alpha(0)
    barras = ax2.bar(d.index, d["Cumulative Return"], color=cores["secondary"], width=0.4, zorder=1)
    ax2.set_ylabel("Retorno Acumulado", fontsize=12)
    ax2.axhline(0, color="gray", linewidth=1, alpha=0.7)
    ax2.yaxis.set_major_formatter(FuncFormatter(lambda y, _: f"{y * 100:.0f}%"))

    ax1.set_xticks(quatro_datas(d.index))
    ax1.tick_params(axis="x", labelsize=14)
    for ax in (ax1, ax2):
        limpar_eixos(ax)

    fig.suptitle("Ibovespa", fontsize=18, fontweight="bold")
    pontos = f"{d['Close'].iloc[-1]:,.0f}".replace(",", ".")
    ax1.legend([linha[0], barras], [f"Pontos Ibovespa: {pontos}",
                                   f"Retorno Acumulado: {d['Cumulative Return'].iloc[-1] * 100:.2f}%"],
               frameon=False, fontsize=12)
    salvar(fig, caminho)


def salvar(fig, caminho):
    fig.tight_layout()
    fig.savefig(caminho, dpi=300, bbox_inches="tight")
    if matplotlib.get_backend().lower() != "agg":
        plt.show()
    plt.close(fig)


# ------------------------------------------------------------------ Word
def substituir(paragrafo, mapa):
    """Troca marcadores mantendo a formatação do trecho quando o marcador está num só trecho."""
    if not any(k in paragrafo.text for k in mapa):
        return
    for run in paragrafo.runs:
        for k, v in mapa.items():
            if k in run.text:
                run.text = run.text.replace(k, v)
    if any(k in paragrafo.text for k in mapa):  # marcador quebrado entre trechos pelo Word
        texto = paragrafo.text
        for k, v in mapa.items():
            texto = texto.replace(k, v)
        paragrafo.text = texto


def montar_word(modelo, graficos, mes_nome, ano, caminho):
    from docx import Document
    from docx.shared import Inches

    doc = Document(modelo)
    mapa = {"{{MONTH}}": mes_nome, "{{YEAR}}": str(ano)}
    paragrafos = list(doc.paragraphs)
    for secao in doc.sections:
        paragrafos += secao.header.paragraphs
    for p in paragrafos:
        substituir(p, mapa)
    for p in doc.paragraphs:
        for marcador, imagem in graficos.items():
            if marcador in p.text:
                p.text = ""
                p.add_run().add_picture(imagem, width=Inches(6))
    doc.save(caminho)


# ------------------------------------------------------------------ principal
def main(mes=MES):
    ano, m = mes_referencia(mes)
    inicio, fim = limites_do_mes(ano, m)
    mes_nome = nome_mes(m)
    saida = pasta_saida(ano, m)
    print(f"Carta mensal de {mes_nome}/{ano} ({inicio:%d/%m} a {fim:%d/%m})")

    df = baixar_indices(inicio, fim)
    gerados = []
    for marca, cfg in MARCAS.items():
        modelo = os.path.join(PASTA_MODELOS, cfg["modelo"])
        if not os.path.exists(modelo):
            print(f"Aviso: modelo da marca {marca} não encontrado em {modelo}. Marca pulada.")
            continue
        graficos = {
            "{{PLOT1}}": os.path.join(saida, f"carta-{marca}-usa-{mes_nome}.png"),
            "{{PLOT2}}": os.path.join(saida, f"carta-{marca}-resumo-{mes_nome}.png"),
            "{{PLOT3}}": os.path.join(saida, f"carta-{marca}-ibov-{mes_nome}.png"),
        }
        grafico_eua(df, cfg["cores"], graficos["{{PLOT1}}"])
        grafico_global(df, cfg["cores"], graficos["{{PLOT2}}"])
        grafico_ibov(df, cfg["cores"], graficos["{{PLOT3}}"])
        docx = os.path.join(saida, f"{cfg['nome_saida']} - {mes_nome} {ano}.docx")
        montar_word(modelo, graficos, mes_nome, ano, docx)
        gerados.append(docx)
        print(f"Pronto: {docx}")
    if not gerados:
        raise RuntimeError(f"Nenhum modelo de carta encontrado em {PASTA_MODELOS}.")
    return gerados


if __name__ == "__main__":
    main()
