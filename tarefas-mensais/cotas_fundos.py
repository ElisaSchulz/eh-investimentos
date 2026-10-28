# -*- coding: utf-8 -*-
"""
Baixa o histórico diário de cota (NAV) dos fundos UCITS da lista fundos.csv pelo
ISIN, usando o pacote gratuito e não oficial `mstarpy` (dados públicos da
Morningstar), e gera uma planilha Excel.

Rodar no Spyder: botão Run (precisa de Python 3.11 ou mais novo). Por padrão
baixa de INICIO até o fim do mês anterior. Para outro mês, preencha MES abaixo.

Todo mês: atualize em fundos.csv as colunas cota_referencia (use ponto como
separador decimal) e data_referencia (AAAA-MM-DD) com as cotas do extrato.
Para incluir ou tirar fundos, edite o mesmo arquivo.

Abas geradas:
    Cotações  : uma linha por data, uma coluna por fundo (NAV na moeda da classe)
    Controle  : compara o NAV baixado com as cotas informadas em fundos.csv
                (o script também confere a moeda: se o NAV na data de referência de uma
                série diferir mais de 3% da sua cota, ele testa as outras séries do mesmo ISIN)
    Notas     : premissas, fonte e limitações
"""

import datetime as dt
import os
import sys
import time
from collections import Counter

sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))

import pandas as pd
from openpyxl import Workbook
from openpyxl.formatting.rule import CellIsRule
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

from comum import ler_csv, limites_do_mes, mes_referencia, nome_mes, pasta_saida

# ------------------------------------------------------------------ parâmetros
MES = None              # None = mês anterior. Ou "AAAA-MM".
INICIO = "2024-01-01"   # primeira data do histórico
REPETIR_NAV = True      # True: repete o NAV do dia anterior nas datas sem cota


def carregar_fundos():
    """Lê fundos.csv. Devolve a lista (nome, ISIN, cota, observação) em ordem alfabética
    e um dicionário ISIN -> data da cota informada."""
    tabela = ler_csv("fundos.csv")
    fundos, datas = [], {}
    for _, f in tabela.iterrows():
        cota = float(f["cota_referencia"].replace(",", ".")) if f["cota_referencia"] else None
        fundos.append((f["nome"], f["isin"].strip().upper(), cota, f["observacao"]))
        if f["data_referencia"]:
            datas[f["isin"].strip().upper()] = dt.date.fromisoformat(f["data_referencia"])
    # Ordem alfabética pelo nome do fundo: define a ordem das colunas e das linhas do Controle
    fundos.sort(key=lambda f: f[0].casefold())
    return fundos, datas


FUNDOS, DATAS_REF = carregar_fundos()

FONTE = "Arial"
AZUL = "0000FF"  # convenção: azul = valor digitado (input)
SHEET_COT = "Cotações"
SHEET_CTL = "Controle"
SHEET_NOTAS = "Notas"


# --------------------------------------------------------------------------
# Download
# --------------------------------------------------------------------------
def _parse_nav(raw):
    """Converte o retorno do mstarpy em pd.Series (índice = data, valor = NAV)."""
    if raw is None:
        return pd.Series(dtype=float)
    df = raw if isinstance(raw, pd.DataFrame) else pd.DataFrame(raw)
    if df.empty:
        return pd.Series(dtype=float)

    cols = {str(c).lower(): c for c in df.columns}
    date_col = next((cols[c] for c in cols if "date" in c), None)
    val_col = cols.get("nav") or cols.get("value") or cols.get("close")
    if date_col is None or val_col is None:
        raise ValueError(f"formato inesperado do mstarpy, colunas: {list(df.columns)}")

    idx = pd.to_datetime(df[date_col], errors="coerce")
    if getattr(idx.dt, "tz", None) is not None:
        idx = idx.dt.tz_localize(None)
    vals = pd.to_numeric(df[val_col], errors="coerce")

    s = pd.Series(vals.values, index=idx.dt.normalize())
    s = s[~s.index.isna()].dropna()
    s = s[~s.index.duplicated(keep="last")].sort_index()
    return s


# Diferença máxima entre o NAV baixado e a cota que você informou para considerar que a
# série está na moeda certa. É larga de propósito (3%): serve para separar moedas diferentes
# (diferenças de 10% ou mais), não para auditar o NAV. A conferência fina fica na aba Controle.
TOL_MOEDA = 0.03


def _candidatos(isin, max_resultados=8):
    """Resultados da busca da Morningstar para o ISIN.

    O mesmo ISIN pode aparecer em mais de uma série (por exemplo, com NAV em moedas diferentes).
    Devolve (lista_de_fundos, confirmado). `confirmado` é False quando nenhum resultado tem
    exatamente o ISIN pedido; nesse caso devolve só o primeiro resultado.
    """
    from mstarpy import Funds  # import tardio para o resto do script ser testável

    todos, vistos = [], set()
    for ir in range(max_resultados):
        try:
            f = Funds(term=isin, itemRange=ir)
        except ValueError:
            break  # acabaram os resultados
        codigo = getattr(f, "code", None) or id(f)
        if codigo in vistos:
            continue
        vistos.add(codigo)
        todos.append(f)
    if not todos:
        raise ValueError("ISIN não encontrado na busca da Morningstar")

    exatos = [f for f in todos if str(getattr(f, "isin", "")).strip().upper() == isin.upper()]
    if exatos:
        return exatos, True
    return todos[:1], False


def _baixar_serie(fundo, start, end):
    raw = fundo.nav(
        dt.datetime.combine(start, dt.time()),
        dt.datetime.combine(end, dt.time()),
        frequency="daily",
    )
    serie = _parse_nav(raw)
    if serie.empty:
        raise ValueError("Morningstar não retornou dados para este ISIN")
    return serie


def _desvio_na_data(serie, ref, check_date):
    """Desvio relativo entre o NAV mais próximo de check_date e a cota informada (ou None)."""
    if ref is None or check_date is None or serie.empty:
        return None
    alvo = pd.Timestamp(check_date)
    perto = serie[(serie.index >= alvo - pd.Timedelta(days=7)) &
                  (serie.index <= alvo + pd.Timedelta(days=3))]
    if perto.empty:
        return None
    pos = abs(perto.index - alvo).argmin()
    return float(perto.iloc[pos]) / ref - 1


def fetch_nav(isin, start, end, ref=None, check_date=None, tentativas=3):
    """Baixa o NAV diário de um ISIN.

    Se `ref` (cota informada em `check_date`) for dada, confere a moeda: percorre as séries
    que a Morningstar devolve para o ISIN e escolhe a primeira cujo NAV nessa data fica a
    menos de TOL_MOEDA da cota informada.

    Devolve (serie, nome_morningstar, isin_confirmado, info) com
    info = {"desvio": float|None, "moeda_ok": bool|None, "indice": int, "total": int}.
    """
    ultimo_erro = None
    for n in range(1, tentativas + 1):
        try:
            candidatos, confirmado = _candidatos(isin)
            melhor = None  # (|desvio|, indice, fundo, serie, desvio)
            escolhido = None
            for k, fundo in enumerate(candidatos):
                serie = _baixar_serie(fundo, start, end)
                desvio = _desvio_na_data(serie, ref, check_date)
                if desvio is None:  # sem referência para comparar: usa esta série
                    escolhido = (k, fundo, serie, None, None)
                    break
                if abs(desvio) <= TOL_MOEDA:
                    escolhido = (k, fundo, serie, desvio, True)
                    break
                if melhor is None or abs(desvio) < melhor[0]:
                    melhor = (abs(desvio), k, fundo, serie, desvio)
            if escolhido is None:  # nenhuma série bateu: fica com a mais próxima e sinaliza
                _, k, fundo, serie, desvio = melhor
                escolhido = (k, fundo, serie, desvio, False)

            k, fundo, serie, desvio, moeda_ok = escolhido
            info = {"desvio": desvio, "moeda_ok": moeda_ok, "indice": k, "total": len(candidatos)}
            return serie, getattr(fundo, "name", "") or "", confirmado, info
        except Exception as e:  # noqa: BLE001 (queremos tentar de novo em qualquer erro)
            ultimo_erro = e
            time.sleep(2 * n)
    raise RuntimeError(str(ultimo_erro))


def baixar_todos(start, end, check_date, pausa=1.5):
    series, status, nomes_ms = {}, {}, {}
    total = len(FUNDOS)
    for i, (nome, isin, cota, _obs) in enumerate(FUNDOS, 1):
        print(f"[{i:02d}/{total}] {isin}  {nome} ... ", end="", flush=True)
        try:
            s, nome_ms, confirmado, info = fetch_nav(isin, start, end, ref=cota,
                                                     check_date=DATAS_REF.get(isin, check_date))
            series[isin] = s
            nomes_ms[isin] = nome_ms

            partes = []
            if info["moeda_ok"] is False:
                status[isin] = (f"ATENÇÃO: NAV difere {info['desvio']:+.1%} da cota informada em todas as "
                                f"{info['total']} série(s) da busca; provável moeda diferente ({len(s)} datas)")
            elif not confirmado:
                status[isin] = f"ATENÇÃO: ISIN não confirmado na busca ({len(s)} datas)"
            else:
                status[isin] = f"OK ({len(s)} datas)"
                if info["indice"] > 0:
                    status[isin] = (f"OK ({len(s)} datas; série #{info['indice'] + 1} de {info['total']}, "
                                    f"escolhida pela moeda)")
            if info["desvio"] is not None:
                partes.append(f"desvio vs cota informada {info['desvio']:+.2%}")
            print(f"{status[isin].split(' (')[0].split(':')[0]}, {len(s)} datas, "
                  f"{s.index.min():%d/%m/%Y} a {s.index.max():%d/%m/%Y}"
                  f"{'; ' + '; '.join(partes) if partes else ''}  [{nome_ms}]")
        except Exception as e:  # noqa: BLE001
            series[isin] = pd.Series(dtype=float)
            nomes_ms[isin] = ""
            status[isin] = f"FALHOU: {str(e)[:120]}"
            print(f"FALHOU ({e})")
        time.sleep(pausa)
    return series, status, nomes_ms


# --------------------------------------------------------------------------
# Planilha
# --------------------------------------------------------------------------
def _f(bold=False, color=None, size=10):
    return Font(name=FONTE, bold=bold, color=color, size=size)


def construir_planilha(series, status, nomes_ms, start, end, check_date, caminho, ffill=True):
    isins = [f[1] for f in FUNDOS]
    df = pd.concat({isin: series[isin] for isin in isins}, axis=1)
    df = df.sort_index()
    df = df[(df.index >= pd.Timestamp(start)) & (df.index <= pd.Timestamp(end))]

    # Repete o NAV do dia anterior quando não há cota na data. Só preenche DEPOIS da primeira
    # cota de cada fundo (antes do lançamento da classe a célula continua vazia).
    if ffill:
        df_com_lacunas = df
        df = df.ffill()
        preenchido = (df_com_lacunas.isna() & df.notna()).values
    else:
        preenchido = pd.DataFrame(False, index=df.index, columns=df.columns).values

    wb = Workbook()

    # ---------------- Aba Cotações ----------------
    ws = wb.active
    ws.title = SHEET_COT
    hdr_fill = PatternFill("solid", start_color="1F3864")
    hdr_font = Font(name=FONTE, bold=True, color="FFFFFF", size=10)
    thin = Side(style="thin", color="BFBFBF")

    ws.cell(row=1, column=1, value="Data")
    ws.cell(row=2, column=1, value="ISIN")
    for j, (nome, isin, _c, _o) in enumerate(FUNDOS, start=2):
        ws.cell(row=1, column=j, value=nome)
        ws.cell(row=2, column=j, value=isin)
    for r in (1, 2):
        for c in range(1, len(FUNDOS) + 2):
            cell = ws.cell(row=r, column=c)
            cell.font = hdr_font
            cell.fill = hdr_fill
            cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    ws.row_dimensions[1].height = 62

    primeira_linha = 3
    fonte_preenchida = Font(name=FONTE, size=10, italic=True, color="808080")
    for pos, (data, linha) in enumerate(df.iterrows()):
        i = primeira_linha + pos
        c = ws.cell(row=i, column=1, value=data.to_pydatetime())
        c.number_format = "dd/mm/yyyy"
        c.font = _f()
        for j, isin in enumerate(isins, start=2):
            v = linha[isin]
            if pd.notna(v):
                cell = ws.cell(row=i, column=j, value=float(v))
                cell.number_format = "#,##0.0000"
                cell.font = fonte_preenchida if preenchido[pos, j - 2] else _f()
    ultima_linha = max(primeira_linha, primeira_linha + len(df) - 1)

    ws.column_dimensions["A"].width = 12
    for j in range(2, len(FUNDOS) + 2):
        ws.column_dimensions[get_column_letter(j)].width = 17
    ws.freeze_panes = "B3"

    # ---------------- Aba Controle ----------------
    wc = wb.create_sheet(SHEET_CTL)
    wc["A1"] = "Conferência das classes: NAV baixado x cotas informadas"
    wc["A1"].font = _f(bold=True, size=12)

    wc["A2"] = "Data de controle"
    wc["B2"] = dt.datetime.combine(check_date, dt.time())
    wc["B2"].number_format = "dd/mm/yyyy"
    wc["A3"] = "Tolerância (diferença máxima aceita)"
    wc["B3"] = 0.005
    wc["B3"].number_format = "0.00%"
    for ref in ("A2", "A3"):
        wc[ref].font = _f(bold=True)
    for ref in ("B2", "B3"):
        wc[ref].font = _f(color=AZUL)
        wc[ref].fill = PatternFill("solid", start_color="FFFF99")
    wc["C2"] = "Altere a data se quiser conferir outro dia (precisa existir na aba Cotações)."
    wc["C3"] = "0,5% é um ponto de partida; diferenças maiores sugerem classe errada ou data diferente."
    for ref in ("C2", "C3"):
        wc[ref].font = _f(color="595959")

    cab = ["Fundo", "ISIN", "Cota informada", "NAV baixado na data", "Dif. %",
           "NAV no dia útil anterior", "Dif. %", "Status", "Download", "Nome na Morningstar", "Observação"]
    linha_cab = 5
    for j, t in enumerate(cab, start=1):
        c = wc.cell(row=linha_cab, column=j, value=t)
        c.font = hdr_font
        c.fill = hdr_fill
        c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    wc.row_dimensions[linha_cab].height = 32

    rng_datas = f"'{SHEET_COT}'!$A${primeira_linha}:$A${ultima_linha}"
    for k, (nome, isin, cota, obs) in enumerate(FUNDOS):
        r = linha_cab + 1 + k
        col = get_column_letter(2 + k)  # coluna do fundo na aba Cotações
        rng_nav = f"'{SHEET_COT}'!${col}${primeira_linha}:${col}${ultima_linha}"
        m = f"MATCH($B$2,{rng_datas},0)"

        wc.cell(row=r, column=1, value=nome).font = _f()
        wc.cell(row=r, column=2, value=isin).font = _f()
        c = wc.cell(row=r, column=3, value=cota)
        c.font = _f(color=AZUL)
        c.number_format = "#,##0.0000"

        wc.cell(row=r, column=4,
                value=f'=IFERROR(IF(INDEX({rng_nav},{m})="","sem dado",INDEX({rng_nav},{m})),"sem dado")')
        wc.cell(row=r, column=5, value=f'=IF(ISNUMBER(D{r}),D{r}/C{r}-1,"")')
        wc.cell(row=r, column=6,
                value=f'=IFERROR(IF(INDEX({rng_nav},{m}-1)="","sem dado",INDEX({rng_nav},{m}-1)),"sem dado")')
        wc.cell(row=r, column=7, value=f'=IF(ISNUMBER(F{r}),F{r}/C{r}-1,"")')
        wc.cell(row=r, column=8,
                value=(f'=IF(AND(NOT(ISNUMBER(E{r})),NOT(ISNUMBER(G{r}))),"sem dado",'
                       f'IF(OR(IF(ISNUMBER(E{r}),ABS(E{r}),9)<=$B$3,'
                       f'IF(ISNUMBER(G{r}),ABS(G{r}),9)<=$B$3),"OK","VERIFICAR"))'))
        wc.cell(row=r, column=9, value=status[isin])
        wc.cell(row=r, column=10, value=nomes_ms.get(isin, ""))
        wc.cell(row=r, column=11, value=obs)

        for cc in (4, 6):
            wc.cell(row=r, column=cc).number_format = "#,##0.0000"
        for cc in (5, 7):
            wc.cell(row=r, column=cc).number_format = "0.00%"
        for cc in range(4, 12):
            wc.cell(row=r, column=cc).font = _f()
        wc.cell(row=r, column=8).alignment = Alignment(horizontal="center")

    ultima_ctl = linha_cab + len(FUNDOS)
    faixa_status = f"H{linha_cab + 1}:H{ultima_ctl}"
    wc.conditional_formatting.add(
        faixa_status, CellIsRule(operator="equal", formula=['"OK"'],
                                 fill=PatternFill("solid", bgColor="C6EFCE"),
                                 font=Font(name=FONTE, color="006100")))
    wc.conditional_formatting.add(
        faixa_status, CellIsRule(operator="equal", formula=['"VERIFICAR"'],
                                 fill=PatternFill("solid", bgColor="FFC7CE"),
                                 font=Font(name=FONTE, color="9C0006")))
    wc.conditional_formatting.add(
        faixa_status, CellIsRule(operator="equal", formula=['"sem dado"'],
                                 fill=PatternFill("solid", bgColor="FFEB9C"),
                                 font=Font(name=FONTE, color="9C5700")))

    for letra, larg in zip("ABCDEFGHIJK", (52, 16, 14, 18, 9, 20, 9, 12, 30, 50, 38)):
        wc.column_dimensions[letra].width = larg
    wc.freeze_panes = f"A{linha_cab + 1}"

    # ---------------- Aba Notas ----------------
    wn = wb.create_sheet(SHEET_NOTAS)
    notas = [
        ("Planilha gerada em", dt.datetime.now().strftime("%d/%m/%Y %H:%M")),
        ("Período solicitado", f"{start:%d/%m/%Y} a {end:%d/%m/%Y}"),
        ("Fonte dos dados", "Morningstar (dados públicos), via pacote não oficial mstarpy."),
        ("Moeda e classe", "NAV na moeda da classe indicada em cada ISIN (todas USD, acumulação). "
                           f"O script compara o NAV de {check_date:%d/%m} com a cota informada e, se a diferença passar de 3%, "
                           "testa as outras séries que a Morningstar devolve para o mesmo ISIN. Se nenhuma bater, "
                           "o status do fundo fica como ATENÇÃO na coluna Download da aba Controle."),
        ("Datas", "A linha de cada data mostra o NAV com a data atribuída pela Morningstar, "
                  "que pode estar 1 dia útil à frente ou atrás da data de valuation do gestor."),
        ("Datas sem cota", ("Quando o fundo não tem NAV na data (feriado local, ou cota ainda não publicada), "
                           "a planilha repete o NAV do dia anterior, em cinza itálico. Antes da primeira cota "
                           "da classe a célula fica vazia. Por causa disso, a aba Controle não acusa 'sem dado' "
                           "nas datas preenchidas." if ffill else
                           "Fundo sem NAV naquela data fica vazio. Nada foi preenchido ou interpolado.")),
        (f"Cotas de {check_date:%d/%m}", "Aba Controle, coluna C: valores informados pela própria usuária (azul)."),
        ("ISINs", "Identificados por pesquisa em sites de dados de mercado (finanzen, Cbonds, "
                  "BlackRock, Barings, fact sheets de gestores). Conferir o Pictet (classe R) no extrato."),
        ("Aviso", "Dados informativos, sem garantia de exatidão. Para qualquer relatório a cliente, "
                  "confirme o NAV no extrato ou no site do gestor."),
    ]
    wn["A1"] = "Notas e premissas"
    wn["A1"].font = _f(bold=True, size=12)
    for i, (k, v) in enumerate(notas, start=3):
        wn.cell(row=i, column=1, value=k).font = _f(bold=True)
        c = wn.cell(row=i, column=2, value=v)
        c.font = _f()
        c.alignment = Alignment(wrap_text=True, vertical="top")
    wn.column_dimensions["A"].width = 22
    wn.column_dimensions["B"].width = 110

    wb.calculation.fullCalcOnLoad = True
    wb.save(caminho)
    return df


# --------------------------------------------------------------------------
def main(mes=MES):
    ano, m = mes_referencia(mes)
    _, end = limites_do_mes(ano, m)
    start = dt.date.fromisoformat(INICIO)
    # Data do Controle: a data de referência mais comum em fundos.csv (ou o fim do mês)
    check = Counter(DATAS_REF.values()).most_common(1)[0][0] if DATAS_REF else end
    caminho = os.path.join(pasta_saida(ano, m), f"cotacoes_fundos_{nome_mes(m).lower()}_{ano}.xlsx")

    print(f"Baixando {len(FUNDOS)} fundos de {start} a {end}...\n")
    series, status, nomes_ms = baixar_todos(start, end, check)
    df = construir_planilha(series, status, nomes_ms, start, end, check, caminho,
                            ffill=REPETIR_NAV)

    print("\n--- Resumo ---")
    falhas = [isin for isin, s in status.items() if s.startswith("FALHOU")]
    atencao = [isin for isin, s in status.items() if s.startswith("ATENÇÃO")]
    if atencao:
        print(f"Aviso: {len(atencao)} fundo(s) com ATENÇÃO (moeda ou ISIN): {', '.join(atencao)}")
        print("Esses não devem ser usados sem conferir. Veja a coluna Download na aba Controle.")
    print(f"Planilha salva em: {caminho}  ({len(df)} datas)")
    if falhas:
        print(f"Aviso: {len(falhas)} fundo(s) falharam: {', '.join(falhas)}")
        print("Rode de novo (às vezes é limite de acesso) ou veja a coluna Download na aba Controle.")
    else:
        print("Todos os fundos baixados. Veja a aba Controle para conferir as classes.")
    return [caminho]


if __name__ == "__main__":
    main()
