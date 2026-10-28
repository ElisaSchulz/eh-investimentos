"""
Preço de fechamento diário de listas de ETFs (Yahoo Finance). Junta os antigos
cotacao_etfs.py e adj_cotacao_etfs.py: cada item de CONJUNTOS gera um arquivo.

    ajustado=False  usa o "Close" bruto (o antigo cotacao_etfs.py)
    ajustado=True   usa o "Adj Close", ajustado por dividendos e splits, que serve
                    para calcular retorno total (o antigo adj_cotacao_etfs.py)

Rodar no Spyder: botão Run. Por padrão vai até o fim do mês anterior.
Para outro mês, preencha MES abaixo (ex.: "2026-09").
"""
import os
import sys
from datetime import timedelta

sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))

import pandas as pd

from comum import fechamentos_yahoo, limites_do_mes, mes_referencia, pasta_saida

# ------------------------------------------------------------------ parâmetros
MES = None          # None = mês anterior. Ou "AAAA-MM".

CONJUNTOS = [
    {"arquivo": "fechamentos", "tickers": ["AGG", "CSPX.L", "FLOA.L", "IVV", "IWDA.L"],
     "inicio": "2024-01-01", "ajustado": False},
    {"arquivo": "fechamentos_ajustados", "tickers": ["AGG"],
     "inicio": "2023-01-01", "ajustado": True},
]

EM_REAIS = False    # True: acrescenta colunas convertidas para reais e o câmbio usado
EXCEL = False       # True: gera .xlsx formatado; False: CSV com ; e vírgula decimal
REPETIR_PRECO = True  # True: feriado numa praça repete o fechamento anterior

# Moeda de negociação por ticker. O que não estiver aqui é tratado como USD.
MOEDAS = {"CSPX.L": "USD", "CSPX.AS": "USD", "CSP1.L": "GBP", "SXR8.DE": "EUR",
          "IVV": "USD", "AGG": "USD"}
EM_PENCE = {"CSP1.L"}  # cotados em pence: o script divide por 100


def moeda_do_ticker(ticker):
    return MOEDAS.get(ticker.upper(), "USD")


def baixar_fechamentos(tickers, inicio, fim_exclusivo, ajustado, preencher=True):
    coluna = "Adj Close" if ajustado else "Close"
    series = {}
    for ticker in tickers:
        fech = fechamentos_yahoo(ticker, inicio, fim_exclusivo, coluna=coluna)
        if fech.empty:
            print(f"Aviso: nenhum dado retornado para {ticker}. Ticker ignorado.")
            continue
        if ticker.upper() in EM_PENCE:
            fech = fech / 100
        series[ticker.upper()] = fech
    if not series:
        raise RuntimeError("Nenhum ticker retornou dados no período informado.")

    df = pd.concat(series, axis=1, sort=False).sort_index()
    df.index.name = "data"
    if preencher:
        faltantes = int(df.isna().sum().sum())
        # ffill só repete valor já conhecido: antes da primeira negociação a célula fica vazia
        df = df.ffill()
        ainda_vazias = int(df.isna().sum().sum())
        if faltantes:
            print(f"{faltantes - ainda_vazias} célula(s) preenchidas com o fechamento anterior.")
        if ainda_vazias:
            print(f"{ainda_vazias} célula(s) sem preço anterior disponível (início da série).")
    return df


def converter_para_brl(df, inicio, fim_exclusivo):
    """Acrescenta uma coluna em reais para cada ticker, usando o câmbio da moeda dele."""
    taxas = {}
    for moeda in {moeda_do_ticker(t) for t in df.columns} - {"BRL"}:
        serie = fechamentos_yahoo(f"{moeda}BRL=X", inicio, fim_exclusivo)
        if serie.empty:
            print(f"Aviso: câmbio {moeda}BRL indisponível. Colunas em {moeda} ficam sem conversão.")
            continue
        taxas[moeda] = serie.reindex(df.index).ffill()  # feriados diferentes entre as praças
    for ticker in list(df.columns):
        if moeda_do_ticker(ticker) in taxas:
            df[f"{ticker}_BRL"] = df[ticker] * taxas[moeda_do_ticker(ticker)]
    for moeda, serie in taxas.items():
        df[f"cambio_{moeda}BRL"] = serie
    return df


def exportar_excel(df, caminho):
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter

    with pd.ExcelWriter(caminho, engine="openpyxl", datetime_format="dd/mm/yyyy") as writer:
        df.to_excel(writer, sheet_name="Fechamento")
        ws = writer.sheets["Fechamento"]
        for celula in ws[1]:
            celula.font = Font(bold=True, color="FFFFFF")
            celula.fill = PatternFill("solid", fgColor="1F3864")
            celula.alignment = Alignment(horizontal="center", vertical="center")
        ws.freeze_panes = "B2"
        for i, nome in enumerate([df.index.name] + list(df.columns), start=1):
            ws.column_dimensions[get_column_letter(i)].width = max(12, len(str(nome)) + 4)
            if i == 1:
                continue
            formato = "0.0000" if str(nome).startswith("cambio_") else "#,##0.00"
            for linha in ws.iter_rows(min_row=2, min_col=i, max_col=i):
                linha[0].number_format = formato


def main(mes=MES):
    ano, m = mes_referencia(mes)
    _, fim = limites_do_mes(ano, m)
    fim_exclusivo = fim + timedelta(days=1)
    saida = pasta_saida(ano, m)
    gerados = []
    for c in CONJUNTOS:
        tipo = "ajustado (Adj Close)" if c["ajustado"] else "bruto (Close)"
        print(f"\n{c['arquivo']}: {', '.join(c['tickers'])}, {tipo}, {c['inicio']} a {fim:%Y-%m-%d}")
        df = baixar_fechamentos(c["tickers"], c["inicio"], fim_exclusivo, c["ajustado"],
                                preencher=REPETIR_PRECO)
        if EM_REAIS:
            df = converter_para_brl(df, c["inicio"], fim_exclusivo)
        if EXCEL:
            caminho = os.path.join(saida, f"{c['arquivo']}.xlsx")
            exportar_excel(df, caminho)
        else:
            caminho = os.path.join(saida, f"{c['arquivo']}.csv")
            df.to_csv(caminho, sep=";", decimal=",", float_format="%.6f", encoding="utf-8-sig")
        print(f"Período: {df.index[0].date()} a {df.index[-1].date()} ({len(df)} datas)")
        print(f"Arquivo salvo: {caminho}")
        print(df.tail(5).to_string())
        gerados.append(caminho)
    return gerados


if __name__ == "__main__":
    main()
